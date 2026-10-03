# -*- coding: utf-8 -*-
"""Excel 报表生成层：把公开批发价格整理成可筛选的市场简报。"""
import pandas as pd


def build_excel(output_path, report_date, sheets):
    """写入多工作表 Excel 日报，并为每张表应用一致的编辑部式版面。"""
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, (rows, note) in sheets.items():
            frame = pd.DataFrame(rows)
            frame.to_excel(writer, sheet_name=sheet_name, index=False, startrow=1)
            _format_sheet(writer, sheet_name, frame, report_date, note)

        if not sheets:
            pd.DataFrame({"提示": ["本次没有抓取到任何数据"]}).to_excel(
                writer, sheet_name="空报表", index=False
            )

    print(f"[完成] 报表已生成：{output_path}")


def _format_sheet(writer, sheet_name, frame, report_date, note):
    """把标题、筛选表头、交替行底色和涨跌提示应用到单张工作表。"""
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    worksheet = writer.sheets[sheet_name]
    navy = "182A38"
    amber = "D9A441"
    ink = "263742"
    muted = "66747A"
    ivory = "FFFDFA"
    stripe = "F5F0E5"
    grid = "E8E1D4"
    positive = "A63E34"  # 国内行情表常用红色表示上涨
    negative = "39745B"  # 绿色表示下跌
    max_col = max(len(frame.columns), 1)
    max_row = len(frame) + 2
    thin_rule = Side(style="thin", color=grid)
    amber_rule = Side(style="medium", color=amber)

    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_properties.tabColor = amber

    title = f"全国农产品价格简报　/　{report_date}"
    if note:
        title += f"　·　{note}"
    for column in range(1, max_col + 1):
        cell = worksheet.cell(row=1, column=column)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.border = Border(bottom=amber_rule)
    title_cell = worksheet.cell(row=1, column=1)
    title_cell.value = title
    title_cell.font = Font(name="微软雅黑", size=14, bold=True, color="FFFFFF")
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    if max_col > 1:
        worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    worksheet.row_dimensions[1].height = 33

    for column, column_name in enumerate(frame.columns, start=1):
        cell = worksheet.cell(row=2, column=column)
        cell.font = Font(name="微软雅黑", size=10, bold=True, color=navy)
        cell.fill = PatternFill("solid", fgColor=amber)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=navy))
    worksheet.row_dimensions[2].height = 27

    for row in range(3, max_row + 1):
        row_fill = ivory if row % 2 else stripe
        worksheet.row_dimensions[row].height = 23
        for column, column_name in enumerate(frame.columns, start=1):
            cell = worksheet.cell(row=row, column=column)
            is_numeric = isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
            cell.font = Font(
                name="微软雅黑", size=10,
                color=ink if is_numeric else muted,
            )
            cell.fill = PatternFill("solid", fgColor=row_fill)
            cell.alignment = Alignment(
                horizontal="right" if is_numeric else "left",
                vertical="center",
            )
            cell.border = Border(bottom=thin_rule)
            if is_numeric:
                if "(%)" in str(column_name) or "涨跌幅" in str(column_name):
                    cell.number_format = '0.00"%"'
                elif "元/公斤" in str(column_name):
                    cell.number_format = "0.00"
                elif "指数" in str(column_name):
                    cell.number_format = "0.00"
                elif isinstance(cell.value, float):
                    cell.number_format = "#,##0.00"

    if len(frame.columns):
        last_column = get_column_letter(len(frame.columns))
        worksheet.auto_filter.ref = f"A2:{last_column}{max_row}"

    for column, column_name in enumerate(frame.columns, start=1):
        longest = len(str(column_name))
        for row in range(3, max_row + 1):
            value = worksheet.cell(row=row, column=column).value
            if value is not None:
                longest = max(
                    longest,
                    sum(2 if ord(char) > 127 else 1 for char in str(value)),
                )
        worksheet.column_dimensions[get_column_letter(column)].width = min(
            max(longest + 3, 12), 38
        )

        if "涨跌幅" in str(column_name) or "涨跌额" in str(column_name):
            column_letter = get_column_letter(column)
            price_range = f"{column_letter}3:{column_letter}{max_row}"
            worksheet.conditional_formatting.add(
                price_range,
                CellIsRule(
                    operator="greaterThan",
                    formula=["0"],
                    fill=PatternFill("solid", fgColor="F8E4DF"),
                    font=Font(color=positive, bold=True),
                ),
            )
            worksheet.conditional_formatting.add(
                price_range,
                CellIsRule(
                    operator="lessThan",
                    formula=["0"],
                    fill=PatternFill("solid", fgColor="E6F0E7"),
                    font=Font(color=negative, bold=True),
                ),
            )

    worksheet.freeze_panes = "A3"
    worksheet.print_title_rows = "1:2"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True


if __name__ == "__main__":
    print("此模块由 main.py 调用，不建议单独运行。")
